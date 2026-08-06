import io
from datetime import date, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Query, Response, HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.db.base import get_db
from app.models.user import User
from app.models.task_update import TaskUpdate
from app.api.v1.endpoints.auth import get_current_user

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

router = APIRouter()


@router.get("/excel")
async def export_excel(
    days: int = Query(default=7, description="Number of days to export"),
    user_id: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    start_date = date.today() - timedelta(days=days)
    
    query = (
        select(TaskUpdate)
        .where(TaskUpdate.report_date >= start_date)
        .options(
            selectinload(TaskUpdate.user),
            selectinload(TaskUpdate.project),
            selectinload(TaskUpdate.task),
        )
    )

    if user.role.name != "ADMIN":
        query = query.where(TaskUpdate.user_id == user.id)
    elif user_id:
        query = query.where(TaskUpdate.user_id == user_id)

    query = query.order_by(TaskUpdate.report_date.desc(), TaskUpdate.created_at.desc())
    result = await db.execute(query)
    updates = result.scalars().all()

    # Create OpenPyXL workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Work Updates"

    # Header styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = [
        "Report Date", "Employee Name", "Project", "Task Title",
        "Morning Planned Work", "Morning Est. Hours", "Blockers",
        "Evening Completed Work", "Hours Logged", "Challenges",
        "Next Day Plan", "Attachment URL"
    ]
    
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    row_idx = 2
    for u in updates:
        row_data = [
            str(u.report_date),
            u.user.full_name if u.user else "N/A",
            u.project.name if u.project else "General",
            u.morning_task_title or (u.task.title if u.task else "N/A"),
            u.morning_planned_work or "N/A",
            float(u.morning_est_hours) if u.morning_est_hours else 0,
            u.morning_blockers or "None",
            u.evening_completed_work or "Pending",
            float(u.evening_hours_spent) if u.evening_hours_spent else 0,
            u.evening_challenges or "None",
            u.evening_next_day_plan or "N/A",
            u.attachment_url or "None"
        ]
        ws.append(row_data)

        # Style data row
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row_idx % 2 == 1:
                cell.fill = alt_fill

        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

    # Adjust column widths
    column_widths = [14, 20, 20, 25, 30, 16, 20, 30, 14, 20, 25, 25]
    for i, col_width in enumerate(column_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = col_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Work_Updates_Report_{days}Days_{date.today().isoformat()}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/ppt")
async def export_ppt(
    days: int = Query(default=7, description="Number of days to export"),
    user_id: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
):
    start_date = date.today() - timedelta(days=days)

    query = (
        select(TaskUpdate)
        .where(TaskUpdate.report_date >= start_date)
        .options(
            selectinload(TaskUpdate.user),
            selectinload(TaskUpdate.project),
            selectinload(TaskUpdate.task),
        )
    )

    if user.role.name != "ADMIN":
        query = query.where(TaskUpdate.user_id == user.id)
    elif user_id:
        query = query.where(TaskUpdate.user_id == user_id)

    query = query.order_by(TaskUpdate.report_date.desc())
    result = await db.execute(query)
    updates = result.scalars().all()

    # Create Presentation deck using python-pptx
    prs = Presentation()
    prs.slide_width = Inches(13.33)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Helper to add dark themed background
    def apply_dark_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate-900

    # Title Slide
    slide = prs.slides.add_slide(blank_slide_layout)
    apply_dark_bg(slide)

    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.33), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"Weekly Work Progress Report"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(59, 130, 246) # Blue-500

    p2 = tf.add_paragraph()
    p2.text = f"Updates for the Last {days} Days ({start_date.isoformat()} to {date.today().isoformat()})"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(148, 163, 184) # Slate-400

    p3 = tf.add_paragraph()
    p3.text = f"Generated for: {user.full_name} ({user.email})"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(203, 213, 225)

    # Summary Slide
    slide2 = prs.slides.add_slide(blank_slide_layout)
    apply_dark_bg(slide2)

    title_box2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.5), Inches(0.8))
    tf2 = title_box2.text_frame
    p = tf2.paragraphs[0]
    p.text = "Executive Performance & Activity Summary"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    # Calculate summary numbers
    total_updates = len(updates)
    total_hours = sum(float(u.evening_hours_spent or 0) for u in updates)
    unique_users = len(set(u.user_id for u in updates))
    unique_projects = len(set(u.project_id for u in updates if u.project_id))

    # Metric Cards
    metrics = [
        ("Total Updates", str(total_updates), RGBColor(59, 130, 246)),
        ("Total Hours Logged", f"{total_hours:.1f} hrs", RGBColor(16, 185, 129)),
        ("Active Members", str(unique_users), RGBColor(168, 85, 247)),
        ("Active Projects", str(unique_projects), RGBColor(245, 158, 11)),
    ]

    card_width = Inches(2.7)
    card_gap = Inches(0.3)
    start_left = Inches(0.8)

    for idx, (label, val, color) in enumerate(metrics):
        left = start_left + idx * (card_width + card_gap)
        card = slide2.shapes.add_shape(1, left, Inches(1.8), card_width, Inches(1.5)) # Rectangle
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(30, 41, 59)
        card.line.color.rgb = color
        card.line.width = Pt(1.5)

        tf_card = card.text_frame
        tf_card.word_wrap = True
        p_val = tf_card.paragraphs[0]
        p_val.text = val
        p_val.font.size = Pt(28)
        p_val.font.bold = True
        p_val.font.color.rgb = color
        p_val.alignment = PP_ALIGN.CENTER

        p_lbl = tf_card.add_paragraph()
        p_lbl.text = label
        p_lbl.font.size = Pt(12)
        p_lbl.font.color.rgb = RGBColor(148, 163, 184)
        p_lbl.alignment = PP_ALIGN.CENTER

    # Detailed Updates Slides (Grouped per day or item)
    # 2 updates per slide layout for clarity
    for idx in range(0, len(updates), 2):
        chunk = updates[idx:idx+2]
        slide_det = prs.slides.add_slide(blank_slide_layout)
        apply_dark_bg(slide_det)

        head_box = slide_det.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.6))
        tf_head = head_box.text_frame
        p_h = tf_head.paragraphs[0]
        p_h.text = f"Daily Updates Detail (Slide {idx // 2 + 1})"
        p_h.font.size = Pt(24)
        p_h.font.bold = True
        p_h.font.color.rgb = RGBColor(255, 255, 255)

        for card_i, u in enumerate(chunk):
            top_pos = Inches(1.2 + card_i * 2.9)
            card = slide_det.shapes.add_shape(1, Inches(0.8), top_pos, Inches(11.73), Inches(2.7))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(30, 41, 59)
            card.line.color.rgb = RGBColor(51, 65, 85)

            tf = card.text_frame
            tf.word_wrap = True

            # Header
            p_top = tf.paragraphs[0]
            p_top.text = f"📅 {u.report_date}  |  👤 {u.user.full_name if u.user else 'N/A'}  |  📁 {u.project.name if u.project else 'General Project'}"
            p_top.font.size = Pt(14)
            p_top.font.bold = True
            p_top.font.color.rgb = RGBColor(59, 130, 246)

            # Planned vs Completed
            p_m = tf.add_paragraph()
            p_m.text = f"• Morning Planned: {u.morning_planned_work or 'N/A'} (Est: {u.morning_est_hours or 0} hrs)"
            p_m.font.size = Pt(12)
            p_m.font.color.rgb = RGBColor(226, 232, 240)

            p_e = tf.add_paragraph()
            p_e.text = f"• Evening Completed: {u.evening_completed_work or 'Pending'} (Actual: {u.evening_hours_spent or 0} hrs)"
            p_e.font.size = Pt(12)
            p_e.font.color.rgb = RGBColor(52, 211, 153)

            if u.morning_blockers or u.evening_challenges:
                p_b = tf.add_paragraph()
                p_b.text = f"• Blockers/Challenges: {u.evening_challenges or u.morning_blockers}"
                p_b.font.size = Pt(11)
                p_b.font.color.rgb = RGBColor(248, 113, 113)

            if u.evening_next_day_plan:
                p_n = tf.add_paragraph()
                p_n.text = f"• Next Day Plan: {u.evening_next_day_plan}"
                p_n.font.size = Pt(11)
                p_n.font.color.rgb = RGBColor(148, 163, 184)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)

    filename = f"Work_Updates_Presentation_{days}Days_{date.today().isoformat()}.pptx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
